#!/usr/bin/env python3
"""Turn real usage into new vocabulary — the loop that makes the web app worth running.

Reads what people actually typed (from the web app's database) and produces the SAME
Excel verify sheet already in use, ranked by how often each spelling appeared:

  1. explicit corrections   — someone said "I typed X, it should be Y"   (highest value)
  2. unmatched spellings    — words the engine could not convert          (the real gap)

Verify the sheet, then fold the rows into `data/vocabulary.csv` exactly like the
hand-curated batches.

Usage:
    PYTHONPATH=src python scripts/export_feedback.py                 # local SQLite
    DATABASE_URL=postgres://... PYTHONPATH=src python scripts/export_feedback.py
    PYTHONPATH=src python scripts/export_feedback.py out.xlsx 200
"""

from __future__ import annotations

import os
import sqlite3
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from sing_khmer_engine.lookup import Engine  # noqa: E402

DATABASE_URL = os.environ.get("DATABASE_URL", "")
SQLITE_PATH = os.environ.get("SQLITE_PATH", "/tmp/sing_khmer.db")


def _rows(sql: str) -> list[tuple]:
    if DATABASE_URL:
        import psycopg

        with psycopg.connect(DATABASE_URL) as conn:
            return conn.cursor().execute(sql).fetchall()
    if not Path(SQLITE_PATH).exists():
        return []
    with sqlite3.connect(SQLITE_PATH) as conn:
        return conn.execute(sql).fetchall()


def collect(eng: Engine, limit: int) -> tuple[list[tuple[str, str, int]], list[tuple[str, int]]]:
    """Return (corrections, unmatched) — both ranked by frequency."""
    corrections = Counter()
    for spelling, khmer in _rows("SELECT spelling, expected_khmer FROM feedback"):
        if spelling:
            corrections[((spelling or "").strip().lower(), (khmer or "").strip())] += 1

    known = set(eng.index)
    unmatched = Counter()
    for (text,) in _rows("SELECT input_text FROM events"):
        for token in (text or "").lower().split():
            core = "".join(ch for ch in token if ch.isalnum())
            # only count words the engine genuinely can't place
            if core and core not in known and not eng.convert(core):
                unmatched[core] += 1

    ranked_c = [(s, k, n) for (s, k), n in corrections.most_common(limit)]
    ranked_u = [(w, n) for w, n in unmatched.most_common(limit) if n >= 1]
    return ranked_c, ranked_u


def write_xlsx(path: Path, corrections, unmatched, eng: Engine) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    hf = PatternFill("solid", fgColor="0F766E")
    hfont = Font(color="FFFFFF", bold=True, size=12)
    thin = Side(style="thin", color="CCCCCC")
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)

    def header(ws, cols):
        for c, h in enumerate(cols, 1):
            x = ws.cell(row=1, column=c, value=h)
            x.fill, x.font = hf, hfont
            x.alignment = Alignment(horizontal="center", vertical="center")
            x.border = bd
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "Corrections from users"
    header(ws, ["#", "They typed", "They say it means", "Times", "✅ / fix"])
    for i, (spelling, khmer, n) in enumerate(corrections, 1):
        ws.cell(row=i + 1, column=1, value=i)
        ws.cell(row=i + 1, column=2, value=spelling).font = Font(name="Consolas", size=12)
        ws.cell(row=i + 1, column=3, value=khmer).font = Font(name="Khmer OS", size=16)
        ws.cell(row=i + 1, column=4, value=n)
    for col, w in zip("ABCDE", [5, 22, 20, 8, 26]):
        ws.column_dimensions[col].width = w

    ws2 = wb.create_sheet("Words we couldn't convert")
    header(ws2, ["#", "They typed", "Times", "Khmer it should be", "note"])
    for i, (word, n) in enumerate(unmatched, 1):
        ws2.cell(row=i + 1, column=1, value=i)
        ws2.cell(row=i + 1, column=2, value=word).font = Font(name="Consolas", size=12)
        ws2.cell(row=i + 1, column=3, value=n)
        ws2.cell(row=i + 1, column=4).font = Font(name="Khmer OS", size=16)
    for col, w in zip("ABCDE", [5, 22, 8, 24, 24]):
        ws2.column_dimensions[col].width = w

    wb.save(path)


def main() -> int:
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "user_feedback.xlsx"
    limit = int(sys.argv[2]) if len(sys.argv) > 2 else 300
    eng = Engine()
    corrections, unmatched = collect(eng, limit)
    source = "postgres" if DATABASE_URL else SQLITE_PATH
    print(f"source: {source}")
    print(f"  corrections from users : {len(corrections)}")
    print(f"  unconvertible spellings: {len(unmatched)}")
    if not corrections and not unmatched:
        print("nothing collected yet — share the app and come back later.")
        return 0
    write_xlsx(out, corrections, unmatched, eng)
    print(f"wrote {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
